"""Local conversion worker. No shell commands, uploads, or in-place writes."""
import csv
import contextlib
import gzip
import html
import importlib.util
import io
import json
import os
import posixpath
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tarfile
import tempfile
import textwrap
import zipfile
import xml.etree.ElementTree as ET
from urllib.parse import unquote

if os.environ.get("CITRUS_LIBRARIES"):
    sys.path.insert(0, os.environ["CITRUS_LIBRARIES"])

IMAGE = [".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".avif", ".heic", ".heif"]
AUDIO = [".mp3", ".m4a", ".wav", ".flac", ".ogg", ".opus", ".aiff", ".wma"]
VIDEO = [".mp4", ".mov", ".mkv", ".webm", ".avi", ".wmv", ".gif", ".m4v"]
TEXT = [".txt", ".md", ".html", ".htm", ".docx", ".epub"]
CREATE_NO_WINDOW = 0x08000000 if os.name == "nt" else 0


def available(name):
    return importlib.util.find_spec(name) is not None


def ffmpeg():
    if os.environ.get("CITRUS_FFMPEG"):
        candidate = Path(os.environ["CITRUS_FFMPEG"])
        return str(candidate) if candidate.is_file() else None
    binary = shutil.which("ffmpeg")
    if not binary and available("imageio_ffmpeg"):
        import imageio_ffmpeg
        binary = imageio_ffmpeg.get_ffmpeg_exe()
    return binary


def office():
    if os.environ.get("CITRUS_OFFICE"):
        candidate = Path(os.environ["CITRUS_OFFICE"])
        return str(candidate) if candidate.is_file() else None
    return shutil.which("soffice") or next((str(p) for p in [
        Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "LibreOffice/program/soffice.exe",
        Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "LibreOffice/program/soffice.exe"
    ] if p.is_file()), None)


def capabilities():
    routes = []
    def add(key, label, extension, sources, warning="", minimum=1, aggregate=False):
        routes.append(dict(Id=key, Label=label, Extension=extension, Sources=sources, Warning=warning, MinimumFiles=minimum, Aggregate=aggregate))
    if available("PIL"):
        from PIL import Image
        if available("pillow_heif"):
            import pillow_heif
            pillow_heif.register_heif_opener()
        Image.init()
        inputs = [ext for ext in IMAGE if ext in Image.registered_extensions()]
        for ext, codec in [("png", "PNG"), ("jpg", "JPEG"), ("bmp", "BMP"), ("tiff", "TIFF"), ("webp", "WEBP"), ("avif", "AVIF"), ("heic", "HEIF")]:
            if codec in Image.SAVE:
                aliases = [".jpg", ".jpeg"] if ext == "jpg" else [".tif", ".tiff"] if ext == "tiff" else [".heic", ".heif"] if ext == "heic" else ["." + ext]
                add("image." + ext, ext.upper(), "." + ext, [e for e in inputs if e not in aliases],
                    "JPG and BMP flatten transparency onto white; multi-frame images are rejected." if ext in ["jpg", "bmp"] else "Multi-frame images are rejected.")
        if available("pymupdf"):
            add("image.pdf", "PDF", ".pdf", inputs)
            add("image.combinepdf", "Combine PDF", ".pdf", inputs, "One image per page; output beside the first selected file.", minimum=2, aggregate=True)
        add("image.resize", "Resize 50%", ".png", inputs, "Creates a half-size PNG. Multi-frame images are rejected.")
        add("image.compress", "Compress JPG", ".jpg", inputs, "Creates a JPEG at quality 70; transparency becomes white.")
    if ffmpeg():
        listing = run_process([ffmpeg(), "-hide_banner", "-encoders"], timeout=15).stdout.decode("utf-8", "replace")
        encoders = {line.split()[1] for line in listing.splitlines() if len(line.split()) > 1 and len(line.split()[0]) == 6 and line.split()[0][0] in "VAS"}
        audio_encoders = {"mp3": "libmp3lame", "m4a": "aac", "wav": "pcm_s16le", "flac": "flac", "ogg": "libvorbis", "opus": "libopus", "aiff": "pcm_s16be", "wma": "wmav2"}
        for ext in ["mp3", "m4a", "wav", "flac", "ogg", "opus", "aiff", "wma"]:
            if audio_encoders[ext] in encoders: add("media." + ext, ext.upper(), "." + ext, [e for e in AUDIO + VIDEO if e != "." + ext])
        for ext in ["mp4", "mov", "mkv", "webm", "avi", "wmv", "gif"]:
            needed = {"webm": {"libvpx-vp9", "libopus"}, "avi": {"mpeg4", "libmp3lame"}, "wmv": {"wmv2", "wmav2"}, "gif": {"gif"}}.get(ext, {"libx264", "aac"})
            if needed <= encoders: add("media." + ext, ext.upper(), "." + ext, [e for e in VIDEO if e != "." + ext])
        if "libx264" in encoders: add("media.noaudio", "Mute video", ".mp4", VIDEO, "Creates an MP4 without audio.")
        if "pcm_s16le" in encoders: add("media.normalize", "Normalize", ".wav", AUDIO, "Normalizes audio to -16 LUFS in a WAV copy.")
    for ext in ["srt", "vtt", "txt"]:
        add("subtitle." + ext, ext.upper(), "." + ext, [e for e in [".srt", ".vtt"] if e != "." + ext],
            "TXT removes timestamps; SRT cannot preserve WebVTT styling and positioning.")
    for ext in ["zip", "tar", "tar.gz"]:
        excluded = [".gz", ".tgz"] if ext == "tar.gz" else ["." + ext]
        add("archive." + ext, ext.upper(), "." + ext, [e for e in [".zip", ".tar", ".gz", ".tgz"] if e not in excluded])
    add("archive.extract", "Extract", ".folder", [".zip", ".tar", ".gz", ".tgz"], "Extracts into a new folder. Links and unsafe paths are rejected.")
    add("archive.packzip", "ZIP archive", ".zip", ["*"], "Packages the selection into one ZIP beside the first item.", aggregate=True)
    if available("pymupdf"):
        add("pdf.compress", "Compress PDF", ".pdf", [".pdf"], "Rewrites and compresses PDF objects; size reduction is not guaranteed.")
        add("pdf.split", "Split pages", ".pdf", [".pdf"], "Creates one PDF per page in a new folder.")
        add("pdf.merge", "Merge PDFs", ".pdf", [".pdf"], "Merges in selection order, saved beside the first selected file.", minimum=2, aggregate=True)
        for ext in ["png", "jpg", "txt", "md", "epub"]:
            add("pdf." + ext, ext.upper() + (" pages" if ext in ["png", "jpg"] else ""), "." + ext, [".pdf"],
                "All pages exported into a new folder at 300 DPI." if ext in ["png", "jpg"] else "Text-layer extraction only; scanned PDFs need OCR.")
        if available("docx") and available("pdf2docx"):
            add("pdf.docx", "Word (layout)", ".docx", [".pdf"], "Reconstructs editable text, tables, images and page layout. Complex layouts and unavailable fonts can differ; scanned pages need OCR.")
        if available("pptx"):
            add("pdf.pptx", "PPTX slides", ".pptx", [".pdf"], "One rendered page per slide. Text is not editable.")
        if available("openpyxl"):
            add("pdf.xlsx", "XLSX tables", ".xlsx", [".pdf"], "Detected tables only. PDF layout is not recreated; scanned tables need OCR.")
        for ext in ["pdf", "txt", "md", "html", "epub"]:
            sources = [e for e in TEXT if e != "." + ext and (e != ".docx" or available("docx")) and (e not in [".epub", ".html", ".htm"] or available("bs4"))]
            add("text." + ext, ext.upper() + (" text" if ext == "pdf" else ""), "." + ext, sources, "Text conversion; complex document layout is not preserved.")
        if available("docx"):
            add("text.docx", "DOCX", ".docx", [e for e in TEXT if e != ".docx" and (e not in [".epub", ".html", ".htm"] or available("bs4"))], "Text paragraphs; complex source layout is not preserved.")
    if available("openpyxl"):
        for ext in ["csv", "tsv", "xlsx"]:
            add("sheet." + ext, ext.upper(), "." + ext, [e for e in [".csv", ".tsv", ".xlsx"] if e != "." + ext],
                "CSV/TSV export uses the first worksheet, cached formula values, and no formatting.")
    if office():
        groups = [([".doc", ".docx", ".odt", ".rtf"], ["pdf", "docx", "odt", "rtf", "txt"]),
                  ([".ppt", ".pptx", ".odp"], ["pdf", "pptx", "odp", "png", "jpg"]),
                  ([".xls", ".xlsx", ".ods", ".csv", ".tsv"], ["pdf", "xlsx", "ods", "csv", "tsv"])]
        for i, (sources, outputs) in enumerate(groups):
            for ext in outputs:
                add("office" + str(i) + "." + ext, ext.upper(), "." + ext, [e for e in sources if e != "." + ext],
                    "Rendered locally by LibreOffice; fonts and layout may differ from Microsoft Office.")
    return routes


def run_process(args, timeout=1800):
    proc = subprocess.run(args, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                          creationflags=CREATE_NO_WINDOW, timeout=timeout, shell=False)
    if proc.returncode:
        raise ValueError(proc.stderr.decode("utf-8", "replace")[-3000:] or "Conversion engine failed.")
    return proc


def image_convert(source, target, output, quality):
    from PIL import Image, ImageOps
    Image.MAX_IMAGE_PIXELS = 80_000_000
    if available("pillow_heif"):
        import pillow_heif
        pillow_heif.register_heif_opener()
    with Image.open(source) as original:
        if getattr(original, "n_frames", 1) > 1:
            raise ValueError("Multi-frame image: choose a media conversion to preserve animation.")
        image = ImageOps.exif_transpose(original)
        if image.width * image.height > 80_000_000:
            raise ValueError("Images above 80 megapixels are not supported.")
        if target == "resize":
            image = image.resize((max(1, image.width // 2), max(1, image.height // 2)))
            target = "png"
        if target == "compress": target, quality = "jpg", 70
        if target in ["jpg", "bmp", "pdf"]:
            rgba = image.convert("RGBA")
            image = Image.new("RGB", image.size, "white")
            image.paste(rgba, mask=rgba.getchannel("A"))
        else:
            image = image.convert("RGBA")
        options = {"quality": quality} if target in ["jpg", "webp", "avif", "heic"] else {}
        image.save(output, format={"jpg": "JPEG", "heic": "HEIF"}.get(target, target.upper()), **options)
        if target != "pdf":
            with Image.open(output) as check:
                check.load()
                if check.size != image.size:
                    raise ValueError("Output dimensions failed validation.")


def media_convert(source, target, output):
    args = [ffmpeg(), "-nostdin", "-hide_banner", "-loglevel", "error", "-protocol_whitelist", "file,pipe", "-i", str(source), "-map_metadata", "-1"]
    audio = {"mp3": ["libmp3lame", "-q:a", "2"], "m4a": ["aac", "-b:a", "192k"], "wav": ["pcm_s16le"],
             "flac": ["flac"], "ogg": ["libvorbis", "-q:a", "5"], "opus": ["libopus", "-b:a", "128k"],
             "aiff": ["pcm_s16be"], "wma": ["wmav2", "-b:a", "192k"]}
    if target == "normalize": args += ["-vn", "-af", "loudnorm=I=-16:TP=-1.5:LRA=11", "-c:a", "pcm_s16le"]
    elif target == "noaudio": args += ["-an", "-c:v", "libx264", "-preset", "fast", "-crf", "23", "-pix_fmt", "yuv420p", "-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2"]
    elif target in audio:
        args += ["-map", "0:a:0", "-vn", "-c:a"] + audio[target]
    elif target == "gif":
        args += ["-an", "-vf", "fps=12,scale=640:-1:flags=lanczos", "-loop", "0"]
    else:
        args += ["-map", "0:v:0", "-map", "0:a:0?", "-vf", "scale=trunc(iw/2)*2:trunc(ih/2)*2", "-pix_fmt", "yuv420p"]
        if target == "webm": args += ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-c:a", "libopus"]
        elif target == "avi": args += ["-c:v", "mpeg4", "-q:v", "4", "-c:a", "libmp3lame"]
        elif target == "wmv": args += ["-c:v", "wmv2", "-b:v", "3M", "-c:a", "wmav2"]
        else:
            args += ["-c:v", "libx264", "-preset", "fast", "-crf", "23", "-c:a", "aac"]
            if target in ["mp4", "mov"]: args += ["-movflags", "+faststart"]
    run_process(args + ["-y", str(output)])
    run_process([ffmpeg(), "-nostdin", "-v", "error", "-i", str(output), "-t", "0.1", "-f", "null", "-"])


def read_text(source):
    if source.stat().st_size > 20 * 1024 * 1024:
        raise ValueError("Text inputs above 20 MB are not supported.")
    data = source.read_bytes()
    if data.startswith((b"\xff\xfe", b"\xfe\xff")): return data.decode("utf-16")
    try: return data.decode("utf-8-sig")
    except UnicodeDecodeError: return data.decode("cp1252")


def subtitle_convert(source, target, output):
    raw = read_text(source).replace("\r\n", "\n").strip()
    blocks = re.split(r"\n\s*\n", raw)
    cues = []
    for block in blocks:
        lines = block.splitlines()
        timing = next((i for i, line in enumerate(lines) if "-->" in line), None)
        if timing is None: continue
        match = re.match(r"\s*((?:\d+:)?\d{2}:\d{2}[.,]\d{3})\s*-->\s*((?:\d+:)?\d{2}:\d{2}[.,]\d{3})", lines[timing])
        if not match: raise ValueError("Invalid subtitle timestamp.")
        times = [t.replace(",", ".") for t in match.groups()]
        times = ["00:" + t if t.count(":") == 1 else t for t in times]
        cues.append((times, "\n".join(lines[timing + 1:])))
    if not cues: raise ValueError("No valid subtitle cues found.")
    if target == "txt": text = "\n\n".join(html.unescape(re.sub(r"<[^>]*>", "", c[1])) for c in cues)
    else:
        text = "WEBVTT\n\n" if target == "vtt" else ""
        for i, (times, content) in enumerate(cues, 1):
            if target == "srt": times = [t.replace(".", ",") for t in times]
            text += (str(i) + "\n" if target == "srt" else "") + " --> ".join(times) + "\n" + content + "\n\n"
    output.write_text(text, encoding="utf-8")


def archive_convert(source, target, output):
    limit = 2 * 1024 ** 3
    with tempfile.TemporaryDirectory(dir=output.parent) as folder:
        root = Path(folder)
        def destination(name):
            path = (root / name).resolve()
            if not path.is_relative_to(root.resolve()): raise ValueError("Unsafe archive path.")
            path.parent.mkdir(parents=True, exist_ok=True)
            return path
        if zipfile.is_zipfile(source):
            with zipfile.ZipFile(source) as archive:
                entries = archive.infolist()
                if len(entries) > 10000 or sum(e.file_size for e in entries) > limit: raise ValueError("Archive expansion limit exceeded.")
                for entry in entries:
                    if (entry.external_attr >> 16) & 0o170000 == 0o120000: raise ValueError("Archive links are not supported.")
                    path = destination(entry.filename)
                    if entry.is_dir(): path.mkdir(exist_ok=True)
                    else:
                        with archive.open(entry) as src, path.open("wb") as dst: shutil.copyfileobj(src, dst)
        elif tarfile.is_tarfile(source):
            with tarfile.open(source) as archive:
                entries = archive.getmembers()
                if len(entries) > 10000 or sum(e.size for e in entries) > limit: raise ValueError("Archive expansion limit exceeded.")
                for entry in entries:
                    path = destination(entry.name)
                    if entry.isdir(): path.mkdir(exist_ok=True)
                    elif entry.isfile():
                        with archive.extractfile(entry) as src, path.open("wb") as dst: shutil.copyfileobj(src, dst)
                    else: raise ValueError("Archive links/devices are not supported.")
        else:
            with gzip.open(source, "rb") as src, destination(source.stem).open("wb") as dst:
                count = 0
                while chunk := src.read(1024 * 1024):
                    count += len(chunk)
                    if count > limit: raise ValueError("Archive expansion limit exceeded.")
                    dst.write(chunk)
        if target == "extract":
            shutil.copytree(root, output)
        elif target == "zip":
            with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
                for path in root.rglob("*"): archive.write(path, path.relative_to(root).as_posix())
        else:
            with tarfile.open(output, "w:gz" if target == "tar.gz" else "w") as archive:
                for path in root.iterdir(): archive.add(path, arcname=path.name)


def extract_text(source):
    ext = source.suffix.lower()
    if ext == ".docx":
        from docx import Document
        doc = Document(source)
        return "\n\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join("\t".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
    if ext == ".epub":
        with zipfile.ZipFile(source) as archive:
            container = ET.fromstring(archive.read("META-INF/container.xml"))
            opf_name = next(node.attrib["full-path"] for node in container.iter() if node.tag.endswith("rootfile"))
            opf = ET.fromstring(archive.read(opf_name))
            items = {node.attrib["id"]: node.attrib.get("href", "") for node in opf.iter() if node.tag.endswith("}item")}
            names = [posixpath.normpath(posixpath.join(posixpath.dirname(opf_name), unquote(items[node.attrib["idref"]])))
                     for node in opf.iter() if node.tag.endswith("}itemref")]
            if any(name.startswith("../") or name.startswith("/") for name in names): raise ValueError("Invalid EPUB chapter path.")
            if len(names) > 10000 or sum(archive.getinfo(n).file_size for n in names) > 20 * 1024 ** 2: raise ValueError("EPUB too large.")
            from bs4 import BeautifulSoup
            return "\n\n".join(BeautifulSoup(archive.read(n), "html.parser").get_text("\n") for n in names)
    text = read_text(source)
    if ext in [".html", ".htm"]:
        from bs4 import BeautifulSoup
        return BeautifulSoup(text, "html.parser").get_text("\n")
    return text


def write_document(text, target, output):
    if target in ["txt", "md"]: output.write_text(text, encoding="utf-8")
    elif target == "html": output.write_text("<!doctype html><meta charset='utf-8'><pre>" + html.escape(text) + "</pre>", encoding="utf-8")
    elif target == "docx":
        from docx import Document
        doc = Document()
        for line in text.splitlines(): doc.add_paragraph(line)
        doc.save(output)
    elif target == "epub":
        with zipfile.ZipFile(output, "w") as archive:
            archive.writestr("mimetype", "application/epub+zip", compress_type=zipfile.ZIP_STORED)
            archive.writestr("META-INF/container.xml", '<?xml version="1.0"?><container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container"><rootfiles><rootfile full-path="book.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')
            archive.writestr("book.opf", '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="id"><metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:identifier id="id">urn:citrus:local</dc:identifier><dc:title>Converted document</dc:title><dc:language>en</dc:language><meta property="dcterms:modified">2026-01-01T00:00:00Z</meta></metadata><manifest><item id="content" href="content.xhtml" media-type="application/xhtml+xml"/><item id="nav" href="nav.xhtml" media-type="application/xhtml+xml" properties="nav"/></manifest><spine><itemref idref="content"/></spine></package>')
            archive.writestr("content.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Document</title></head><body><pre>' + html.escape(text) + '</pre></body></html>')
            archive.writestr("nav.xhtml", '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops"><head><title>Contents</title></head><body><nav epub:type="toc"><ol><li><a href="content.xhtml">Document</a></li></ol></nav></body></html>')
    elif target == "pdf":
        import pymupdf as fitz
        doc = fitz.open()
        lines = [wrapped for line in text.splitlines() for wrapped in (textwrap.wrap(line, 88, replace_whitespace=False) or [""])] or [""]
        for start in range(0, len(lines), 48):
            page = doc.new_page()
            page.insert_htmlbox(fitz.Rect(40, 40, 555, 800), "<div>" + html.escape("\n".join(lines[start:start + 48])).replace("\n", "<br/>") + "</div>", css="* {font-family: sans-serif; font-size: 10pt;}")
        doc.save(output); doc.close()


def pdf_convert(source, target, output):
    import pymupdf as fitz
    with fitz.open(source) as doc:
        if doc.needs_pass: raise ValueError("Password-protected PDFs are not supported.")
        if target == "compress":
            doc.save(output, garbage=4, deflate=True)
        elif target == "split":
            output = output.with_suffix(".pages"); output.mkdir()
            for index in range(len(doc)):
                with fitz.open() as part:
                    part.insert_pdf(doc, from_page=index, to_page=index)
                    part.save(output / ("page-" + str(index + 1).zfill(4) + ".pdf"))
        elif target in ["png", "jpg"]:
            output = output.with_suffix(".pages")
            output.mkdir()
            for index, page in enumerate(doc):
                if page.rect.width * page.rect.height * (300 / 72) ** 2 > 80_000_000: raise ValueError("PDF page too large.")
                page.get_pixmap(dpi=300, alpha=False).save(output / ("page-" + str(index + 1).zfill(4) + "." + target))
            return output
        elif target == "docx":
            # Do not silently discard scan-only pages or publish a partially parsed document.
            for index, page in enumerate(doc):
                if not page.get_text().strip() and page.get_images():
                    raise ValueError("Page " + str(index + 1) + " has no PDF text layer. OCR is required for scanned pages before editable Word conversion.")
            from pdf2docx import Converter as LayoutConverter
            from docx import Document
            converter = LayoutConverter(str(source))
            try:
                settings = converter.default_settings
                settings.update(ignore_page_error=False, raw_exceptions=True)
                converter.parse(**settings)
                converted = Document()
                for page in converter.pages:
                    if not page.finalized:
                        raise ValueError("Word layout parsing did not complete for page " + str(page.id + 1) + ".")
                    page.make_docx(converted)
                    # pdf2docx decides whether to start a page by checking body
                    # paragraphs. An empty/table-only first page needs an anchor.
                    if not converted.paragraphs:
                        converted.add_paragraph()
                converted.save(output)
            finally:
                converter.close()
            converted = Document(output)
            if len(doc) and not converted.sections:
                raise ValueError("Word conversion produced no document sections.")
        elif target == "xlsx":
            import openpyxl
            book = openpyxl.Workbook()
            book.remove(book.active)
            for page_number, page in enumerate(doc, 1):
                for table_number, table in enumerate(page.find_tables().tables, 1):
                    sheet = book.create_sheet("Page" + str(page_number) + "_Table" + str(table_number))
                    for row in table.extract():
                        sheet.append(row)
                        for cell in sheet[sheet.max_row]:
                            if isinstance(cell.value, str): cell.data_type = "s"
            if not book.worksheets: raise ValueError("No tables detected. Scanned tables need OCR.")
            book.save(output)
        elif target == "pptx":
            from pptx import Presentation
            from pptx.util import Inches
            presentation = Presentation()
            if len(doc):
                presentation.slide_width = Inches(doc[0].rect.width / 72)
                presentation.slide_height = Inches(doc[0].rect.height / 72)
            for page in doc:
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                slide.shapes.add_picture(io.BytesIO(page.get_pixmap(dpi=120).tobytes("png")), 0, 0,
                                         presentation.slide_width, presentation.slide_height)
            presentation.save(output)
        else:
            text = "\n\n".join(page.get_text(sort=True) for page in doc)
            if not text.strip(): raise ValueError("No PDF text layer. OCR is required for scanned documents.")
            write_document(text, target, output)
    return output


def sheet_convert(source, target, output):
    import openpyxl
    if source.suffix.lower() == ".xlsx":
        book = openpyxl.load_workbook(source, read_only=True, data_only=True)
        try: rows = [list(row) for row in book.worksheets[0].iter_rows(values_only=True)]
        finally: book.close()
    else:
        rows = list(csv.reader(io.StringIO(read_text(source)), delimiter="\t" if source.suffix.lower() == ".tsv" else ","))
    if target == "xlsx":
        book = openpyxl.Workbook()
        for row in rows:
            book.active.append(row)
            # Imported text is data, never executable spreadsheet formulas.
            for cell in book.active[book.active.max_row]:
                if isinstance(cell.value, str): cell.data_type = "s"
        book.save(output)
    else:
        with output.open("w", encoding="utf-8", newline="") as stream:
            csv.writer(stream, delimiter="\t" if target == "tsv" else ",").writerows(rows)


def convert(request):
    source, output = Path(request["Source"]).resolve(), Path(request["Output"]).resolve()
    route = next((route for route in capabilities() if route["Id"] == request["Route"]), None)
    if not route or ("*" not in route["Sources"] and source.suffix.lower() not in route["Sources"]):
        raise ValueError("This route is unavailable for the input file.")
    sources = [Path(os.path.abspath(path)) for path in (request.get("Sources") or [request["Source"]])]
    if len(sources) < route["MinimumFiles"] or ("*" not in route["Sources"] and any(path.suffix.lower() not in route["Sources"] for path in sources)):
        raise ValueError("The selection does not support this operation.")
    family, target = request["Route"].split(".", 1)
    if target == "packzip" and family == "archive":
        total = count = 0
        with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
            for source_item in sources:
                paths = [source_item]
                while paths:
                    path = paths.pop()
                    if path.resolve() == output.parent: continue  # Never package our own staging directory.
                    if path.is_symlink() or path.is_junction(): raise ValueError("Archive links/junctions are not supported.")
                    count += 1
                    if path.is_file(): total += path.stat().st_size
                    if count > 10000 or total > 2 * 1024 ** 3: raise ValueError("Archive size/entry limit exceeded.")
                    archive.write(path, path.relative_to(source_item.parent).as_posix())
                    if path.is_dir():
                        for child in path.iterdir():
                            paths.append(child)
                            if count + len(paths) > 10000: raise ValueError("Archive entry limit exceeded.")
    elif route["Aggregate"]:
        import pymupdf as fitz
        with fitz.open() as combined:
            for path in sources:
                if family == "image":
                    temporary = output.parent / "image-page.pdf"
                    image_convert(path, "pdf", temporary, request.get("Quality", 90))
                    path = temporary
                with fitz.open(path) as document:
                    if document.needs_pass: raise ValueError("Password-protected PDFs are not supported.")
                    combined.insert_pdf(document)
            combined.save(output)
    elif family == "image": image_convert(source, target, output, request.get("Quality", 90))
    elif family == "media": media_convert(source, target, output)
    elif family == "subtitle": subtitle_convert(source, target, output)
    elif family == "archive": archive_convert(source, target, output)
    elif family == "pdf": output = pdf_convert(source, target, output)
    elif family == "text": write_document(extract_text(source), target, output)
    elif family == "sheet": sheet_convert(source, target, output)
    elif family.startswith("office"):
        with tempfile.TemporaryDirectory(dir=output.parent) as folder:
            profile_path = Path(folder) / "profile"
            (profile_path / "user").mkdir(parents=True)
            (profile_path / "user/registrymodifications.xcu").write_text('<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>', encoding="utf-8")
            profile = profile_path.as_uri()
            conversion = "pdf" if target in ["png", "jpg"] else target
            if family == "office2" and target in ["csv", "tsv"]:
                conversion = target + ":Text - txt - csv (StarCalc):" + ("9" if target == "tsv" else "44") + ",34,76,1"
            run_process([office(), "-env:UserInstallation=" + profile, "--headless", "--norestore", "--convert-to", conversion, "--outdir", folder, str(source)], timeout=180)
            result = Path(folder) / (source.stem + "." + ("pdf" if target in ["png", "jpg"] else target))
            if not result.is_file(): raise ValueError("LibreOffice did not produce the requested format.")
            if target in ["png", "jpg"]: output = pdf_convert(result, target, output)
            else: shutil.copyfile(result, output)
    if not output.exists() or (output.is_file() and output.stat().st_size == 0):
        raise ValueError("Conversion produced no output.")
    return output


if __name__ == "__main__":
    try:
        if sys.argv[1] == "--capabilities":
            with contextlib.redirect_stdout(sys.stderr): result = capabilities()
            print(json.dumps(result))
        else:
            request = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8-sig"))
            with contextlib.redirect_stdout(sys.stderr): result = convert(request)
            print(json.dumps({"Output": str(result)}))
    except Exception as error:
        print(json.dumps({"Error": str(error)}))
        sys.exit(1)
