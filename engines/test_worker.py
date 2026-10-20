"""Integration checks using real local codecs and generated, non-private fixtures."""
import hashlib
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
import zipfile
sys.path.insert(0, str(Path(__file__).parent))
import worker


class EngineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="CitrusEngineTests-")
        cls.root = Path(cls.temp.name)
        cls.fixtures = {}
        from PIL import Image
        for ext, fmt in [("jpg", "JPEG"), ("png", "PNG"), ("bmp", "BMP"), ("tiff", "TIFF"), ("webp", "WEBP")]:
            path = cls.root / ("sample." + ext)
            Image.new("RGB", (64, 48), "coral").save(path, fmt)
            cls.fixtures["." + ext] = path
        for ext in ["txt", "md", "html"]:
            path = cls.root / ("sample." + ext)
            path.write_text("Citrus engine test\nSecond line.", encoding="utf-8")
            cls.fixtures["." + ext] = path
        for ext, content in [("srt", "1\n00:00:01,000 --> 00:00:02,000\nCitrus\n"),
                             ("vtt", "WEBVTT\n\n00:00:01.000 --> 00:00:02.000\nCitrus\n"),
                             ("csv", 'name,value\n"comma, name",42\n'), ("tsv", "name\tvalue\nCitrus\t42\n")]:
            path = cls.root / ("sample." + ext); path.write_text(content, encoding="utf-8"); cls.fixtures["." + ext] = path
        path = cls.root / "sample.rtf"; path.write_text(r"{\rtf1\ansi Citrus engine test}", encoding="ascii"); cls.fixtures[".rtf"] = path
        if worker.available("fitz"):
            import pymupdf as fitz
            path = cls.root / "sample.pdf"
            with fitz.open() as doc:
                for _ in range(2):
                    page = doc.new_page(); page.insert_text((72, 72), "Citrus engine test")
                    page.draw_rect(fitz.Rect(50, 100, 250, 160)); page.draw_line((150, 100), (150, 160)); page.draw_line((50, 130), (250, 130))
                    for x, y, text in [(55, 120, "Name"), (155, 120, "Value"), (55, 150, "Citrus"), (155, 150, "42")]: page.insert_text((x, y), text)
                doc.save(path)
            cls.fixtures[".pdf"] = path
        if worker.available("docx"):
            from docx import Document
            path = cls.root / "sample.docx"; doc = Document(); doc.add_paragraph("Citrus engine test"); doc.save(path); cls.fixtures[".docx"] = path
        if worker.available("openpyxl"):
            import openpyxl
            path = cls.root / "sample.xlsx"; book = openpyxl.Workbook(); book.active.append(["Citrus", 42]); book.save(path); cls.fixtures[".xlsx"] = path
        if worker.available("pptx"):
            from pptx import Presentation
            path = cls.root / "sample.pptx"; presentation = Presentation(); presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = "Citrus"; presentation.save(path); cls.fixtures[".pptx"] = path
            if worker.office():
                profile = (cls.root / "office-profile").as_uri()
                worker.run_process([worker.office(), "-env:UserInstallation=" + profile, "--headless", "--convert-to", "odp", "--outdir", str(cls.root), str(path)])
                cls.fixtures[".odp"] = cls.root / "sample.odp"
        path = cls.root / "sample.zip"
        with zipfile.ZipFile(path, "w") as archive: archive.writestr("folder/hello.txt", "Citrus")
        cls.fixtures[".zip"] = path
        path = cls.root / "sample.tar"
        import tarfile
        with tarfile.open(path, "w") as archive: archive.add(cls.fixtures[".txt"], arcname="hello.txt")
        cls.fixtures[".tar"] = path
        if worker.ffmpeg():
            wav = cls.root / "sample.wav"
            worker.run_process([worker.ffmpeg(), "-nostdin", "-v", "error", "-f", "lavfi", "-i", "sine=frequency=440:duration=0.3", "-y", str(wav)])
            cls.fixtures[".wav"] = wav
            mp4 = cls.root / "sample.mp4"
            worker.run_process([worker.ffmpeg(), "-nostdin", "-v", "error", "-f", "lavfi", "-i", "color=c=orange:s=64x48:r=12:d=0.3", "-i", str(wav), "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest", "-y", str(mp4)])
            cls.fixtures[".mp4"] = mp4
            mkv = cls.root / "sample.mkv"
            worker.run_process([worker.ffmpeg(), "-nostdin", "-v", "error", "-i", str(mp4), "-c", "copy", "-y", str(mkv)])
            cls.fixtures[".mkv"] = mkv

    @classmethod
    def tearDownClass(cls): cls.temp.cleanup()

    def test_every_advertised_profile(self):
        routes = worker.capabilities()
        self.assertGreater(len(routes), 30)
        for index, route in enumerate(routes):
            with self.subTest(route=route["Id"]):
                source = self.fixtures[".txt"] if "*" in route["Sources"] else next((self.fixtures[ext] for ext in route["Sources"] if ext in self.fixtures), None)
                self.assertIsNotNone(source, route["Id"] + " needs a fixture")
                before = hashlib.sha256(source.read_bytes()).digest()
                output = self.root / ("output-" + str(index) + route["Extension"])
                result = worker.convert({"Source": str(source), "Sources": [str(source)] * route["MinimumFiles"], "Output": str(output), "Route": route["Id"], "Quality": 90})
                self.assertTrue(result.exists())
                self.assertEqual(before, hashlib.sha256(source.read_bytes()).digest())
                if result.is_file(): self.assertGreater(result.stat().st_size, 0)
                if route["Id"] in ["pdf.png", "pdf.jpg", "pdf.split"]: self.assertEqual(len(list(result.iterdir())), 2)
                if route["Id"] in ["pdf.merge", "image.combinepdf"]:
                    import pymupdf as fitz
                    with fitz.open(result) as doc: self.assertEqual(len(doc), 4 if route["Id"] == "pdf.merge" else 2)
                if route["Id"] == "pdf.docx":
                    from docx import Document
                    self.assertIn("Citrus", "\n".join(p.text for p in Document(result).paragraphs))
                if route["Id"] == "pdf.xlsx":
                    import openpyxl
                    book = openpyxl.load_workbook(result)
                    try: self.assertEqual(book.worksheets[0]["A2"].value, "Citrus")
                    finally: book.close()
                if route["Id"] == "pdf.pptx":
                    from pptx import Presentation
                    self.assertEqual(len(Presentation(result).slides), 2)
                if route["Id"] == "text.pdf":
                    import pymupdf as fitz
                    with fitz.open(result) as doc: self.assertIn("Citrus", doc[0].get_text())

    def test_archive_traversal_rejected(self):
        path = self.root / "unsafe.zip"
        with zipfile.ZipFile(path, "w") as archive: archive.writestr("../escaped.txt", "bad")
        with self.assertRaisesRegex(ValueError, "Unsafe"):
            worker.archive_convert(path, "tar", self.root / "unsafe.tar")
        self.assertFalse((self.root / "escaped.txt").exists())

    def test_pdf_word_layout(self):
        if not worker.available("pdf2docx"):
            self.skipTest("Layout engine unavailable")
        import io
        import pymupdf as fitz
        from PIL import Image
        from docx import Document
        source = self.root / "formatted.pdf"
        image = io.BytesIO()
        Image.new("RGB", (80, 40), "orange").save(image, "PNG")
        with fitz.open() as pdf:
            page = pdf.new_page(width=595, height=842)
            page.insert_text((60, 80), "Layout heading", fontname="hebo", fontsize=22, color=(0, 0, 1))
            page.insert_text((60, 120), "Body in Times", fontname="tiro", fontsize=13)
            page.insert_text((60, 150), "Italic emphasis", fontname="heit", fontsize=13)
            for x in [60, 240, 420]: page.draw_line((x, 190), (x, 270))
            for y in [190, 230, 270]: page.draw_line((60, y), (420, y))
            for x, y, text in [(70, 215, "Product"), (250, 215, "Price"), (70, 255, "Citrus"), (250, 255, "42")]:
                page.insert_text((x, y), text, fontsize=12)
            page.insert_image(fitz.Rect(60, 310, 180, 370), stream=image.getvalue())
            pdf.new_page(width=595, height=842)  # Intentional blank page must survive.
            page = pdf.new_page(width=842, height=595)
            page.insert_text((60, 80), "Landscape final page", fontname="tiro", fontsize=16)
            pdf.save(source)
        original = hashlib.sha256(source.read_bytes()).hexdigest()
        result = worker.pdf_convert(source, "docx", self.root / "formatted.docx")
        document = Document(result)
        runs = [run for paragraph in document.paragraphs for run in paragraph.runs]
        heading = next(run for run in runs if "Layout heading" in run.text)
        self.assertAlmostEqual(heading.font.size.pt, 22, delta=0.5)
        self.assertTrue(heading.bold)
        self.assertEqual(str(heading.font.color.rgb), "0000FF")
        body = next(run for run in runs if "Body in Times" in run.text)
        self.assertIn("Times", body.font.name)
        self.assertAlmostEqual(body.font.size.pt, 13, delta=0.5)
        self.assertTrue(next(run for run in runs if "Italic emphasis" in run.text).italic)
        self.assertTrue(any(paragraph.paragraph_format.space_before or paragraph.paragraph_format.line_spacing
                            for paragraph in document.paragraphs))
        self.assertTrue(any([[cell.text for cell in row.cells] for row in table.rows] == [["Product", "Price"], ["Citrus", "42"]]
                            for table in document.tables), "PDF table must become editable Word cells")
        self.assertGreaterEqual(len(document.part._element.xpath(".//w:drawing")), 1)
        self.assertGreaterEqual(len(document.sections), 3)
        self.assertAlmostEqual(document.sections[0].page_width.pt, 595, delta=1)
        self.assertAlmostEqual(document.sections[-1].page_width.pt, 842, delta=1)
        self.assertEqual(original, hashlib.sha256(source.read_bytes()).hexdigest())
        if worker.office():
            rendered = self.root / "layout-render"; rendered.mkdir(exist_ok=True)
            profile = (self.root / "layout-office-profile").as_uri()
            worker.run_process([worker.office(), "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf", "--outdir", str(rendered), str(result)])
            with fitz.open(rendered / "formatted.pdf") as pdf:
                self.assertEqual(len(pdf), 3, "Rendered Word file must retain the fixture's three pages")
                self.assertIn("Layout heading", pdf[0].get_text())
                self.assertFalse(pdf[1].get_text().strip())
                self.assertIn("Landscape final page", pdf[2].get_text())
                self.assertAlmostEqual(pdf[2].rect.width, 842, delta=2)

    def test_pdf_word_leading_trailing_blank_pages(self):
        if not worker.available("pdf2docx"):
            self.skipTest("Layout engine unavailable")
        import pymupdf as fitz
        from docx import Document
        source = self.root / "blank-edges.pdf"
        with fitz.open() as pdf:
            pdf.new_page()
            page = pdf.new_page(); page.insert_text((72, 72), "Only the middle page has text")
            pdf.new_page(); pdf.save(source)
        result = worker.pdf_convert(source, "docx", self.root / "blank-edges.docx")
        self.assertEqual(len(Document(result).sections), 3)
        if worker.office():
            rendered = self.root / "blank-render"; rendered.mkdir(exist_ok=True)
            profile = (self.root / "blank-office-profile").as_uri()
            worker.run_process([worker.office(), "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf", "--outdir", str(rendered), str(result)])
            with fitz.open(rendered / "blank-edges.pdf") as pdf:
                self.assertEqual(len(pdf), 3)
                self.assertFalse(pdf[0].get_text().strip())
                self.assertIn("Only the middle", pdf[1].get_text())
                self.assertFalse(pdf[2].get_text().strip())

    def test_pdf_word_scan_rejected(self):
        if not worker.available("pdf2docx"):
            self.skipTest("Layout engine unavailable")
        import io
        import pymupdf as fitz
        from PIL import Image
        source = self.root / "scan-word.pdf"
        image = io.BytesIO(); Image.new("RGB", (100, 100), "white").save(image, "PNG")
        with fitz.open() as pdf:
            page = pdf.new_page(); page.insert_image(page.rect, stream=image.getvalue()); pdf.save(source)
        output = self.root / "scan-word.docx"
        with self.assertRaisesRegex(ValueError, "OCR"):
            worker.pdf_convert(source, "docx", output)
        self.assertFalse(output.exists())

    def test_cli_stdout_is_only_json(self):
        import subprocess
        request = self.root / "request.json"
        request.write_text(json.dumps({"Source": str(self.fixtures[".pdf"]), "Output": str(self.root / "cli.xlsx"), "Route": "pdf.xlsx", "Quality": 90}), encoding="utf-8")
        result = subprocess.run([sys.executable, str(Path(worker.__file__)), str(request)], capture_output=True, text=True, timeout=30, creationflags=worker.CREATE_NO_WINDOW)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertTrue(Path(json.loads(result.stdout)["Output"]).is_file())

    def test_subtitle_timing(self):
        output = self.root / "timing.vtt"
        worker.subtitle_convert(self.fixtures[".srt"], "vtt", output)
        self.assertIn("00:00:01.000 --> 00:00:02.000", output.read_text())

    def test_folder_pack_and_pdf_merge_order(self):
        folder = self.root / "picked-folder"; folder.mkdir()
        (folder / "child.txt").write_text("Citrus", encoding="utf-8")
        output = self.root / "folder.zip"
        worker.convert({"Source": str(folder), "Sources": [str(folder)], "Output": str(output), "Route": "archive.packzip", "Quality": 90})
        with zipfile.ZipFile(output) as archive: self.assertEqual(archive.read("picked-folder/child.txt"), b"Citrus")
        staging = folder / ".citrus-work-test"; staging.mkdir()
        inside = staging / "result.zip"
        worker.convert({"Source": str(folder), "Sources": [str(folder)], "Output": str(inside), "Route": "archive.packzip", "Quality": 90})
        with zipfile.ZipFile(inside) as archive: self.assertFalse(any(".citrus-work-test" in name for name in archive.namelist()))
        import pymupdf as fitz
        second = self.root / "second.pdf"
        with fitz.open() as doc: doc.new_page().insert_text((72, 72), "Second document"); doc.save(second)
        merged = self.root / "merged.pdf"
        worker.convert({"Source": str(self.fixtures[".pdf"]), "Sources": [str(self.fixtures[".pdf"]), str(second)], "Output": str(merged), "Route": "pdf.merge", "Quality": 90})
        with fitz.open(merged) as doc:
            self.assertEqual(len(doc), 3)
            self.assertIn("Citrus", doc[0].get_text()); self.assertIn("Second", doc[-1].get_text())

    def test_csv_cells_do_not_become_formulas(self):
        import openpyxl
        source = self.root / "formula.csv"; source.write_text("=1+1\n", encoding="utf-8")
        output = self.root / "formula.xlsx"; worker.sheet_convert(source, "xlsx", output)
        book = openpyxl.load_workbook(output)
        try: self.assertEqual(book.active["A1"].data_type, "s")
        finally: book.close()

    def test_scanned_pdf_needs_ocr(self):
        import pymupdf as fitz
        source = self.root / "scan.pdf"
        with fitz.open() as doc: doc.new_page(); doc.save(source)
        with self.assertRaisesRegex(ValueError, "OCR"):
            worker.pdf_convert(source, "txt", self.root / "scan.txt")


if __name__ == "__main__": unittest.main(verbosity=2)
